#!/usr/bin/env python3
"""Convert a Branch summary export (CSV/Excel) into BigQuery-ready CSV.

Usage:
  python3 bigquery/convert_branch_export.py ~/Downloads/summary-table-export.csv
  python3 bigquery/convert_branch_export.py ~/Downloads/export.xlsx

Output:
  bigquery/data/branch_upload_YYYYMMDD_HHMMSS.csv
"""

from __future__ import annotations

import csv
import re
import sys
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "data"

MONEY_RE = re.compile(r"[\$,\s]")

OUT_COLS = [
    "date",
    "ad_partner",
    "campaign",
    "platform",
    "ad_partner_3p",
    "clicks",
    "installs",
    "register",
    "complete_registration",
    "initiate_purchases",
    "purchases",
    "cost",
    "revenue",
    "ecpi",
    "cpp",
    "ecpc",
    "rc_trial_cancelled_event",
    "rc_expiration_event",
    "rc_cancellation_event",
    "rc_trial_started_event",
    "rc_product_change_event",
    "loaded_at",
]


def parse_money(v):
    if v is None:
        return None
    s = str(v).strip()
    if s == "" or s.lower() in ("null", "none"):
        return None
    s = MONEY_RE.sub("", s)
    if s == "":
        return None
    try:
        return str(Decimal(s))
    except InvalidOperation:
        return None


def parse_int(v):
    if v is None:
        return None
    s = str(v).strip()
    if s == "" or s.lower() in ("null", "none"):
        return None
    s = MONEY_RE.sub("", s)
    if s == "":
        return None
    try:
        return str(int(Decimal(s)))
    except Exception:
        return None


def parse_date(v):
    if v is None:
        return None
    if hasattr(v, "date"):
        try:
            return v.date().isoformat()
        except Exception:
            pass
    s = str(v).strip()
    for fmt in ("%m/%d/%y", "%m/%d/%Y", "%Y/%m/%d", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def norm_key(k: str) -> str:
    return re.sub(r"\s+", " ", str(k).strip().lower())


FIELD_MAP = {
    # Branch summary export names
    "date": "date",
    "ad partner": "ad_partner",
    "campaign": "campaign",
    "platform": "platform",
    "ad partner (3p)": "ad_partner_3p",
    "clicks": "clicks",
    "installs": "installs",
    "register": "register",
    "complete registration": "complete_registration",
    "complete_registration": "complete_registration",
    "initiate purchases": "initiate_purchases",
    "initiate_purchase": "initiate_purchases",
    "initiate purchase": "initiate_purchases",
    "initiate_purchases": "initiate_purchases",
    "purchases": "purchases",
    "purchase": "purchases",
    "cost": "cost",
    "revenue": "revenue",
    "cpi": "ecpi",
    "ecpi": "ecpi",
    "cpp": "cpp",
    "ecpc": "ecpc",
    "rc_trial_cancelled_event": "rc_trial_cancelled_event",
    "rc_expiration_event": "rc_expiration_event",
    "rc_cancellation_event": "rc_cancellation_event",
    "rc_trial_started_event": "rc_trial_started_event",
    "rc_product_change_event": "rc_product_change_event",
    # Already-normalized upload CSV names
    "ad_partner": "ad_partner",
    "ad_partner_3p": "ad_partner_3p",
}

# Required destination fields → label shown in error messages (Branch export names)
REQUIRED_FIELDS: list[tuple[str, str]] = [
    ("date", "date"),
    ("ad_partner", "ad partner"),
    ("campaign", "campaign"),
    ("platform", "platform"),
    ("ad_partner_3p", "ad partner (3p)"),
    ("clicks", "clicks"),
    ("installs", "installs"),
    ("register", "REGISTER"),
    ("complete_registration", "COMPLETE_REGISTRATION"),
    ("initiate_purchases", "INITIATE_PURCHASE"),
    ("purchases", "PURCHASE"),
    ("cost", "cost"),
    ("revenue", "revenue"),
    ("ecpi", "eCPI"),
    ("ecpc", "eCPC"),
    ("rc_trial_cancelled_event", "rc_trial_cancelled_event"),
    ("rc_expiration_event", "rc_expiration_event"),
    ("rc_cancellation_event", "rc_cancellation_event"),
    ("rc_trial_started_event", "rc_trial_started_event"),
    ("rc_product_change_event", "rc_product_change_event"),
]


def read_rows(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xls"}:
        try:
            import openpyxl  # type: ignore
        except ImportError as e:
            raise SystemExit(
                "Excel needs openpyxl. Run: python3 -m pip install openpyxl\n"
                f"Original error: {e}"
            )
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        raw = [[c if c is not None else "" for c in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
        # Find header row that includes a "date" column (Branch Summary exports)
        header_idx = None
        for i, row in enumerate(raw):
            if not row:
                continue
            keys = [norm_key(c) for c in row if c not in ("", None)]
            if "date" in keys:
                header_idx = i
                break
        if header_idx is None:
            raise SystemExit(
                "Could not find a header row with a 'date' column in Excel. "
                "In Branch, Compare by must include date before export."
            )
        headers = [str(h) if h is not None else "" for h in raw[header_idx]]
        rows = []
        for row in raw[header_idx + 1 :]:
            if not any(cell not in ("", None) for cell in row):
                continue
            item = {}
            for h, v in zip(headers, row):
                item[h] = v
            rows.append(item)
        return rows

    # CSV / Branch export with metadata preamble
    text = path.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    date_header_idx = None
    fallback_header_idx = None
    for i, line in enumerate(lines):
        cols = [norm_key(h) for h in next(csv.reader([line])) if h.strip()]
        if not cols:
            continue
        if "date" in cols:
            date_header_idx = i
            break
        if fallback_header_idx is None and "ad partner" in cols and "campaign" in cols:
            fallback_header_idx = i
    header_idx = date_header_idx if date_header_idx is not None else fallback_header_idx
    if header_idx is None:
        return list(csv.DictReader(lines))
    return list(csv.DictReader(lines[header_idx:]))


def read_header_keys(path: Path) -> set[str]:
    """Normalized header names present in the export (lowercase)."""
    rows = read_rows(path)
    if rows:
        return {norm_key(k) for k in rows[0].keys() if k is not None and str(k).strip()}

    # Header-only file: re-open just for field names
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xls"}:
        import openpyxl  # type: ignore

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            keys = [norm_key(c) for c in row if c not in ("", None)]
            if "date" in keys:
                wb.close()
                return set(keys)
        wb.close()
        return set()

    text = path.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    for line in lines:
        if "date" in {norm_key(h) for h in next(csv.reader([line]))}:
            return {norm_key(h) for h in next(csv.reader([line])) if h.strip()}
    return set()


def mapped_fields_from_headers(header_keys: set[str]) -> set[str]:
    present: set[str] = set()
    for src, dest in FIELD_MAP.items():
        if src in header_keys:
            present.add(dest)
    return present


def missing_required_columns(path: Path) -> list[str]:
    """Return Branch-style labels for required columns missing from the export."""
    header_keys = read_header_keys(path)
    present = mapped_fields_from_headers(header_keys)
    missing = []
    for dest, label in REQUIRED_FIELDS:
        if dest not in present:
            missing.append(label)
    return missing


def convert_to_rows(path: Path) -> tuple[list[dict[str, Any]], int]:
    """Return (normalized rows, skipped_count) from a Branch summary export."""
    src_rows = read_rows(path)
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
    out_rows: list[dict[str, Any]] = []
    skipped = 0

    for r in src_rows:
        keyed = {norm_key(k): v for k, v in r.items() if k is not None}
        mapped: dict[str, Any] = {}
        for src, dest in FIELD_MAP.items():
            if src in keyed and dest not in mapped:
                mapped[dest] = keyed[src]

        d = parse_date(mapped.get("date") or keyed.get("date"))
        if not d:
            skipped += 1
            continue

        out_rows.append(
            {
                "date": d,
                "ad_partner": (str(mapped.get("ad_partner") or "").strip() or None),
                "campaign": (str(mapped.get("campaign") or "").strip() or None),
                "platform": (str(mapped.get("platform") or "").strip() or None),
                "ad_partner_3p": (str(mapped.get("ad_partner_3p") or "").strip() or None),
                "clicks": parse_int(mapped.get("clicks")),
                "installs": parse_int(mapped.get("installs")),
                "register": parse_int(mapped.get("register")),
                "complete_registration": parse_int(mapped.get("complete_registration")),
                "initiate_purchases": parse_int(mapped.get("initiate_purchases")),
                "purchases": parse_int(mapped.get("purchases")),
                "cost": parse_money(mapped.get("cost")),
                "revenue": parse_money(mapped.get("revenue")),
                "ecpi": parse_money(mapped.get("ecpi")),
                "cpp": parse_money(mapped.get("cpp")),
                "ecpc": parse_money(mapped.get("ecpc")),
                "rc_trial_cancelled_event": parse_int(mapped.get("rc_trial_cancelled_event")),
                "rc_expiration_event": parse_int(mapped.get("rc_expiration_event")),
                "rc_cancellation_event": parse_int(mapped.get("rc_cancellation_event")),
                "rc_trial_started_event": parse_int(mapped.get("rc_trial_started_event")),
                "rc_product_change_event": parse_int(mapped.get("rc_product_change_event")),
                "loaded_at": now,
            }
        )
    return out_rows, skipped


def convert(path: Path) -> Path:
    out_rows, skipped = convert_to_rows(path)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = OUT_DIR / f"branch_upload_{stamp}.csv"
    with out_path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=OUT_COLS)
        w.writeheader()
        w.writerows(out_rows)

    dates = [r["date"] for r in out_rows]
    print(f"rows: {len(out_rows)}  skipped: {skipped}")
    if dates:
        print(f"date range: {min(dates)} -> {max(dates)}")
    print(f"output: {out_path}")
    return out_path


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        raise SystemExit(1)
    path = Path(sys.argv[1]).expanduser().resolve()
    if not path.exists():
        raise SystemExit(f"File not found: {path}")
    convert(path)


if __name__ == "__main__":
    main()
